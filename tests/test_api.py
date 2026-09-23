from __future__ import annotations

import io
import hashlib
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app
from app.storage import Store
from app.worker import run_once


DEMO = Path(__file__).resolve().parents[1] / "REPLAN_demo_inputs.xlsx"


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("REPLAN_DATA_DIR", str(tmp_path))
    monkeypatch.setenv("REPLAN_DEMO_TOKEN", "test-token")
    monkeypatch.delenv("API_KEY", raising=False)
    return TestClient(app)


def request(client, method, path, **kwargs):
    response = getattr(client, method)(path, headers={"Authorization": "Bearer test-token"}, **kwargs)
    assert response.status_code < 400, response.text
    return response.json()


def baseline(client):
    project_id = request(client, "post", "/api/projects", json={"mode": "REPLAY"})["project_id"]
    preview = request(client, "post", f"/api/projects/{project_id}/imports", files={"file": ("demo.xlsx", DEMO.read_bytes())})
    assert len(preview["tasks"]) == 20
    confirmed = request(client, "post", f"/api/projects/{project_id}/imports/{preview['import_id']}/confirm", json={})
    return project_id, preview, confirmed


def test_original_upload_is_preserved(client):
    project_id = request(client, "post", "/api/projects", json={"mode": "REPLAY"})["project_id"]
    content = DEMO.read_bytes()
    preview = request(client, "post", f"/api/projects/{project_id}/imports", files={"file": ("demo.xlsx", content)})
    assert preview["_upload"]["sha256"] == hashlib.sha256(content).hexdigest()
    response = client.get(
        f"/api/projects/{project_id}/imports/{preview['import_id']}/original",
        headers={"Authorization": "Bearer test-token"},
    )
    assert response.status_code == 200 and response.content == content


def test_watch_plan_rejects_unapproved_source_host(client):
    project_id = request(client, "post", "/api/projects", json={"mode": "LIVE"})["project_id"]
    response = client.put(
        f"/api/projects/{project_id}/watch-plan",
        headers={"Authorization": "Bearer test-token"},
        json={"enabled": True, "source_allowlist": ["https://127.0.0.1/private"]},
    )
    assert response.status_code == 422


def e01(client, project_id, preview):
    event = request(
        client, "post", f"/api/projects/{project_id}/events",
        json={
            "event_id": "E01", "content": preview["events"][0]["body"],
            "mode": "REPLAY", "data_origin": "SYNTHETIC",
            "simulation_as_of": "2026-09-23T09:00:00+02:00",
        },
    )
    return event["event_id"]


def test_e01_budget_replan_approval_commit_and_export(client):
    project_id, preview, baseline_version = baseline(client)
    event_id = e01(client, project_id, preview)
    queued = request(client, "post", f"/api/projects/{project_id}/analyses", json={"event_id": event_id})
    assert run_once(Store())
    result = request(client, "get", f"/api/runs/{queued['run_id']}")
    assert result["run"]["status"] == "succeeded"
    assert result["run"]["data"]["agent_status"] == "llm_unavailable"
    scenarios = {tuple(item["data"]["option_ids"]): item for item in result["scenarios"]}
    assert scenarios[()]["data"]["finish_date"] == "2026-10-30"
    assert scenarios[("OPT-02",)]["data"]["finish_date"] == "2026-10-29"
    assert scenarios[("OPT-03",)]["data"]["finish_date"] == "2026-10-28"
    assert not scenarios[("OPT-03",)]["data"]["budget_met"]
    assert not any(item["data"]["target_met"] and item["data"]["budget_met"] for item in result["scenarios"])

    new_run = request(client, "post", f"/api/runs/{queued['run_id']}/replan", json={"budget_krw": 6000000})
    assert run_once(Store())
    updated = request(client, "get", f"/api/runs/{new_run['run_id']}")
    option = next(item for item in updated["scenarios"] if item["data"]["option_ids"] == ["OPT-03"])
    assert option["data"]["budget_met"] and option["data"]["target_met"]
    scenario_id = option["id"]

    unapproved = client.post(f"/api/scenarios/{scenario_id}/commit", headers={"Authorization": "Bearer test-token"})
    assert unapproved.status_code == 409
    prepared = request(client, "post", f"/api/scenarios/{scenario_id}/prepare")
    assert len(prepared["actions"]) == 1
    pending = client.post(
        f"/api/scenarios/{scenario_id}/approve", headers={"Authorization": "Bearer test-token"},
        json={"actor": "demo-admin", "confirmed_conditions": option["data"]["required_confirmations"]},
    )
    assert pending.status_code == 409
    request(client, "patch", f"/api/actions/{prepared['actions'][0]['id']}", json={"state": "ACCEPTED"})
    request(client, "post", f"/api/scenarios/{scenario_id}/approve", json={"actor": "demo-admin", "confirmed_conditions": option["data"]["required_confirmations"]})
    request(client, "patch", f"/api/actions/{prepared['actions'][0]['id']}", json={"state": "REJECTED"})
    revoked = client.post(f"/api/scenarios/{scenario_id}/commit", headers={"Authorization": "Bearer test-token"})
    assert revoked.status_code == 409
    request(client, "patch", f"/api/actions/{prepared['actions'][0]['id']}", json={"state": "ACCEPTED"})
    committed = request(client, "post", f"/api/scenarios/{scenario_id}/commit")
    assert committed["version_id"] != baseline_version["version_id"]
    assert request(client, "post", f"/api/scenarios/{scenario_id}/commit")["duplicate"]

    response = client.get(f"/api/projects/{project_id}/export?version_id={committed['version_id']}", headers={"Authorization": "Bearer test-token"})
    assert response.status_code == 200
    book = load_workbook(io.BytesIO(response.content), read_only=True)
    assert book.active["D1"].value == committed["version_hash"]
    task_rows = {row[0]: row for row in book.active.iter_rows(min_row=3, values_only=True)}
    assert task_rows["T20"][5] == "2026-10-28"


def test_duplicate_event_and_stale_scenario(client):
    project_id, preview, _ = baseline(client)
    event_id = e01(client, project_id, preview)
    repeated = request(client, "post", f"/api/projects/{project_id}/events", json={"event_id": "E01", "content": preview["events"][0]["body"], "mode": "REPLAY", "data_origin": "SYNTHETIC", "simulation_as_of": "2026-09-23T09:00:00+02:00"})
    assert repeated["duplicate"] and repeated["event_id"] == event_id
    queued = request(client, "post", f"/api/projects/{project_id}/analyses", json={"event_id": event_id, "budget_krw": 6000000})
    assert run_once(Store())
    scenarios = request(client, "get", f"/api/runs/{queued['run_id']}")["scenarios"]
    first = next(item for item in scenarios if item["data"]["option_ids"] == ["OPT-03"])
    second = next(item for item in scenarios if item["data"]["option_ids"] == ["OPT-02"])
    for item in (first, second):
        actions = request(client, "post", f"/api/scenarios/{item['id']}/prepare")["actions"]
        for action in actions:
            request(client, "patch", f"/api/actions/{action['id']}", json={"state": "ACCEPTED"})
        request(client, "post", f"/api/scenarios/{item['id']}/approve", json={"actor": "demo-admin", "confirmed_conditions": item["data"]["required_confirmations"]})
    request(client, "post", f"/api/scenarios/{first['id']}/commit")
    conflict = client.post(f"/api/scenarios/{second['id']}/commit", headers={"Authorization": "Bearer test-token"})
    assert conflict.status_code == 409


def test_revised_excel_requires_diff_confirmation(client):
    project_id, _, _ = baseline(client)
    workbook = load_workbook(DEMO)
    workbook["일정표"]["E10"] = "2026-09-30"
    changed = io.BytesIO()
    workbook.save(changed)
    preview = request(client, "post", f"/api/projects/{project_id}/imports", files={"file": ("revised.xlsx", changed.getvalue())})
    assert preview["import_kind"] == "change"
    assert preview["diff"]["summary"]["changed"] >= 1
    confirmed = request(client, "post", f"/api/projects/{project_id}/imports/{preview['import_id']}/confirm", json={})
    assert confirmed["event"]["channel"] == "revised_excel"
    assert "T03" in confirmed["event"]["related_task_ids"]


def test_demo_auth_required(client):
    assert client.get("/api/projects").status_code == 401
